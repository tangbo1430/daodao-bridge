from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"F:\IM\docs\华为端与唐僧叨叨互通_叨叨研发网络对接说明.docx")
BLUE, GREEN, ORANGE, GRAY = "143C7D", "198754", "D97706", "5E6A78"
LIGHT_BLUE, LIGHT_GREEN, LIGHT_ORANGE, LIGHT_GRAY = "EAF2FB", "EAF7F0", "FFF4E5", "F3F5F7"


def font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def text(p, value, size=10.5, bold=False, color=None):
    r = p.add_run(value)
    font(r, size, bold, color)
    return r


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(cell, color="D5DCE5", size="5"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = b.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            b.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:color"), color)


def margins(cell, value=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = tcPr.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar")
        tcPr.append(m)
    for side in ("top", "start", "bottom", "end"):
        e = m.find(qn("w:" + side))
        if e is None:
            e = OxmlElement("w:" + side)
            m.append(e)
        e.set(qn("w:w"), str(value))
        e.set(qn("w:type"), "dxa")


def heading(doc, value, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    text(p, value, 15 if level == 1 else 12, True, BLUE if level == 1 else "2F6FB3")
    return p


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    text(p, value)


def callout(doc, title, value, fill=LIGHT_BLUE, accent="2F6FB3"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill); borders(c, accent, "9"); margins(c, 140)
    p = c.paragraphs[0]
    text(p, title, 10.5, True, accent)
    p2 = c.add_paragraph(); text(p2, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, BLUE); borders(c, "FFFFFF", "4"); margins(c)
        c.width = Cm(widths[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; text(p, h, 9.3, True, "FFFFFF")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]
            shade(c, "FFFFFF" if ri % 2 == 0 else LIGHT_GRAY); borders(c); margins(c)
            c.width = Cm(widths[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); text(p, str(value), 9.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
s = doc.sections[0]; s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.5); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
doc.styles["Normal"].font.name = "Microsoft YaHei"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
doc.styles["Normal"].font.size = Pt(10.5)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(18)
text(p, "华为端与唐僧叨叨互通", 22, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; text(p, "叨叨研发网络对接说明", 16, True, "2F6FB3")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; text(p, "V1.1  ·  2026-09-28", 10, False, GRAY)

callout(doc, "一句话要求", "请为独立 Bridge 提供一条受控的私网访问通道，使其能访问叨叨业务 HTTP 和悟空 IM WebSocket。Bridge 不必与叨叨同机，叨叨服务端和悟空核心不改代码。", LIGHT_GREEN, GREEN)

heading(doc, "1. 为什么需要私网通道")
for v in [
    "A+ 方案为每个华为用户保持一条悟空 PC(2) WebSocket，预计 300–500 条长期连接。",
    "公网代理、NAT 或 WAF 的空闲超时可能批量断开 WebSocket，造成消息延迟和补偿压力。",
    "私网/VPN 能限制访问来源、降低公网暴露面，并便于排查连接和延迟问题。",
    "这不是要求物理运营商专线；只要是稳定、受控、支持长连接的私网通道即可。",
]: bullet(doc, v)

heading(doc, "2. 请提供的网络条件")
table(doc, ["Bridge 访问目标", "协议/端口", "用途"], [
    ("叨叨业务 HTTP（master 或内网 LB）", "TCP 8090；如有 TLS 则 HTTPS", "注册、登录、同步、好友、群和媒体接口"),
    ("悟空 IM WebSocket（master 或内网 LB）", "TCP 15200；跨公网必须 WSS", "PC(2) connect、recv、recvack、send、ping"),
    ("腾讯 PHP → Bridge", "HTTPS 8443（双方约定）", "腾讯回调转交；由腾讯侧配置，不需要叨叨回连"),
], [6.3, 4.1, 6.0])

callout(doc, "推荐落地方式", "Bridge 与叨叨在同一云网络时，配置同 VPC/对等网络和安全组；不在同一云时，配置站点到站点 IPsec VPN。只需放行 Bridge 固定源 IP 到 8090、15200。", LIGHT_ORANGE, ORANGE)

heading(doc, "3. 叨叨研发需要回填的信息")
for v in [
    "业务 HTTP 的真实内网 URL、Host、完整路径和 token 请求头规则。",
    "悟空 WebSocket 的真实 URL、路径、是否 WSS，以及 PC(2) connect 鉴权规则。",
    "Bridge 所在网段可使用的路由、DNS、白名单和安全组配置。",
    "测试租户、邀请码、测试账号、Token 有效期和限频。",
    "脱敏的 connect、recv、recvack、send、ping 实际请求和成功响应。",
]: bullet(doc, v)

heading(doc, "4. 最小安全要求")
for v in [
    "只放行 Bridge 固定源 IP 到确认后的 8090、15200；不开放叨叨 MySQL、Redis 或管理端口。",
    "跨公网必须使用 HTTPS/WSS；私网内可使用受控的 HTTP/WS。",
    "Token、密码、邀请码和密钥通过环境变量或密钥文件提供，不写入代码和普通日志。",
    "Bridge 使用独立 MySQL/Redis，不连接叨叨数据库。",
]: bullet(doc, v)

heading(doc, "5. 验收方式")
for v in [
    "Bridge-1、Bridge-2 都能访问 8090 和 15200。",
    "测试 PC(2) 连接保持 30 分钟，心跳正常，没有周期性断线。",
    "完成一次 connect、recv、recvack、send，脱敏帧进入契约测试。",
    "断网恢复后能自动重连并补齐消息，主备切换不重复投递。",
    "叨叨服务端与悟空 IM 代码、配置和数据库保持不变。",
]: bullet(doc, v)

callout(doc, "如果只能提供公网地址", "公网 HTTPS + WSS 也可以作为联调或正式方案，但必须提供稳定的 WebSocket 路径、TLS、IP 白名单、长连接超时和 500 条连接容量；只有公网 HTTP、没有可用 WSS 时，不能完成 A+ 实时链路。", LIGHT_BLUE, "2F6FB3")

for sec in doc.sections:
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text(f, "叨叨研发网络对接说明  |  ", 8.5, False, GRAY)
    r = f.add_run(); b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " PAGE "; e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.extend([b, i, e])

doc.save(OUT)
print(OUT)
