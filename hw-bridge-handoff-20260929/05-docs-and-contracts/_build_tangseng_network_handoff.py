from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"F:\IM\docs\华为端与唐僧叨叨互通_叨叨研发网络对接说明.docx")

BLUE = "143C7D"
MID_BLUE = "2F6FB3"
LIGHT_BLUE = "EAF2FB"
GREEN = "198754"
LIGHT_GREEN = "EAF7F0"
ORANGE = "D97706"
LIGHT_ORANGE = "FFF4E5"
GRAY = "5E6A78"
LIGHT_GRAY = "F3F5F7"
WHITE = "FFFFFF"
RED = "B42318"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D5DCE5", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(p, text, *, bold=False, color=None, size=10.5):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return r


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else MID_BLUE)
    return p


def add_bullet(doc, text, level=0, color=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    add_text(p, text, color=color)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    add_text(p, text)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE, accent=MID_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(17)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, "10")
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_text(p, title, bold=True, color=accent, size=11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_text(p2, body, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, "FFFFFF", "4")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, bold=True, color=WHITE, size=9.5)
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, str(value), size=9.2)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "华为端与唐僧叨叨互通 · 网络对接说明  |  ", color=GRAY, size=8.5)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(54)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "华为端与唐僧叨叨互通", bold=True, color=BLUE, size=25)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_before = Pt(4)
add_text(p2, "叨叨研发网络对接说明", bold=True, color=MID_BLUE, size=20)
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_before = Pt(18)
add_text(p3, "A+ 架构 · Bridge 私网接入", color=GREEN, size=13)

cover_table = doc.add_table(rows=4, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_table.autofit = False
for row in cover_table.rows:
    for cell in row.cells:
        set_cell_border(cell, "DDE3EA", "4")
        set_cell_margins(cell, 100, 160, 100, 160)
cover_rows = [
    ("文档版本", "V1.0"),
    ("编制日期", "2026-09-28"),
    ("对接对象", "唐僧叨叨研发 / 运维 / 网络负责人"),
    ("实施边界", "叨叨服务端与悟空 IM 零代码、零配置修改"),
]
for i, (k, v) in enumerate(cover_rows):
    set_cell_shading(cover_table.cell(i, 0), LIGHT_BLUE)
    add_text(cover_table.cell(i, 0).paragraphs[0], k, bold=True, color=BLUE, size=9.5)
    add_text(cover_table.cell(i, 1).paragraphs[0], v, size=9.5)
    cover_table.cell(i, 0).width = Cm(4)
    cover_table.cell(i, 1).width = Cm(11.5)

doc.add_paragraph()
add_callout(
    doc,
    "结论先行",
    "不要求 Bridge 与叨叨服务部署在同一台机器，也不要求采购物理专线。需要的是一条受控、稳定、支持长期 TCP/WebSocket 连接的私网通道，使独立 Bridge 能访问叨叨业务 HTTP 和悟空 IM WebSocket。优先采用同 VPC、VPC 对等连接或站点到站点 IPsec VPN。",
    LIGHT_GREEN,
    GREEN,
)

doc.add_page_break()

add_heading(doc, "1. 文档目的", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
add_text(p, "本文件用于叨叨研发和网络团队配置 A+ 互通所需的网络条件。Bridge 是独立消息中转层，通过叨叨现有 HTTP 与悟空 IM WebSocket 接口工作，不进入叨叨进程、不读取叨叨数据库，也不要求修改叨叨服务端代码或悟空 IM 核心。")

add_callout(
    doc,
    "网络目标",
    "让 Bridge 主备实例能够主动发起到叨叨业务入口和悟空入口的连接。叨叨侧无需主动回连 Bridge；只需允许已登记的 Bridge 源地址访问指定目标端口。",
)

add_heading(doc, "2. 目标拓扑", 1)
topo = doc.add_table(rows=3, cols=5)
topo.alignment = WD_TABLE_ALIGNMENT.CENTER
topo.autofit = False
labels = [
    (0, 0, "腾讯 PHP\n业务后端", LIGHT_BLUE, BLUE),
    (0, 1, "HTTPS/HMAC\n→", WHITE, GRAY),
    (0, 2, "Bridge-1 / Bridge-2\n独立 Go 服务", LIGHT_ORANGE, ORANGE),
    (0, 3, "私网通道\n→", WHITE, GRAY),
    (0, 4, "叨叨业务 HTTP\nmaster:8090", LIGHT_GREEN, GREEN),
    (1, 2, "独立 MySQL + Redis\n不使用叨叨数据库", LIGHT_GRAY, GRAY),
    (2, 2, "PC(2) 代理连接池\n每个华为用户 1 条", LIGHT_ORANGE, ORANGE),
    (2, 3, "WebSocket\n→", WHITE, GRAY),
    (2, 4, "悟空 IM WebSocket\nmaster:15200", LIGHT_GREEN, GREEN),
]
for row in topo.rows:
    for cell in row.cells:
        set_cell_border(cell, "D5DCE5", "5")
        set_cell_margins(cell, 120, 100, 120, 100)
        cell.width = Cm(3.2)
for r, c, label, fill, color in labels:
    cell = topo.cell(r, c)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, line in enumerate(label.split("\n")):
        if index:
            p.add_run("\n")
        add_text(p, line, bold=index == 0, color=color, size=9.5)
doc.add_paragraph()

add_heading(doc, "3. 叨叨侧需要提供的网络通道", 1)
p = doc.add_paragraph()
add_text(p, "推荐提供三层私网互通，Bridge 通过固定源 IP 访问叨叨侧固定内网域名或负载均衡地址。这里的“专线”指逻辑上的受控私网通道，不强制要求物理运营商专线。")

add_table(
    doc,
    ["优先级", "方式", "适用条件", "叨叨侧工作", "评价"],
    [
        ("1", "同 VPC / 同内网", "Bridge 可部署在叨叨所在云或机房", "分配子网、路由和安全组", "最简单、最稳定"),
        ("2", "VPC 对等连接 / 云企业网", "双方 VPC 可建立私网互通", "配置路由表和安全组", "适合正式生产"),
        ("3", "站点到站点 IPsec VPN", "Bridge 与叨叨分属不同云或机房", "提供 VPN 网关、对端网段和预共享参数", "无需改应用，推荐"),
        ("4", "公网 TLS + IP 白名单", "HTTP 与 WebSocket 都能公网稳定提供", "提供 HTTPS/WSS、白名单和限流", "可联调或作为备选"),
        ("5", "物理专线", "有合规、跨境或高 SLA 要求", "运营商线路和路由接入", "当前 300–500 用户通常无必要"),
    ],
    [1.2, 3.4, 4.2, 4.2, 3.3],
)

add_heading(doc, "4. 必须开放的目标", 1)
add_table(
    doc,
    ["源", "目标", "协议", "方向", "用途", "是否需要叨叨回连"],
    [
        ("Bridge-1/2 固定 IP", "master:8090 或实际内网 LB", "TCP/HTTP 或 HTTPS", "Bridge → 叨叨", "注册、登录、关系、群、同步、媒体接口", "否"),
        ("Bridge-1/2 固定 IP", "master:15200 或实际内网 LB", "TCP/WebSocket；跨公网须 WSS", "Bridge → 悟空", "PC(2) connect、recv、recvack、send、ping", "否"),
        ("腾讯 PHP 出口 IP", "Bridge:8443（或双方约定端口）", "HTTPS + HMAC", "腾讯 PHP → Bridge", "转发腾讯之后回调和业务指令", "不涉及叨叨"),
    ],
    [2.7, 3.7, 3.5, 2.4, 4.2, 2.3],
)

add_callout(
    doc,
    "最小权限原则",
    "不开放 MySQL、Redis、SSH、管理后台或悟空节点内部端口。只允许 Bridge 固定源地址访问确认后的业务 HTTP 和 WebSocket 入口。",
    LIGHT_ORANGE,
    ORANGE,
)

add_heading(doc, "5. 为什么建议私网通道", 1)
for text in [
    "长期连接：A+ 为每个已启用的华为用户保持一条悟空 PC(2) WebSocket。预计 300–500 条常驻连接，公网 NAT、代理或 WAF 的空闲超时可能造成周期性断线。",
    "实时性：目标是叨叨 → 华为文本 P95 小于 1 秒、华为 → 叨叨小于 1.5 秒。私网链路能减少跨公网抖动和重连带来的补拉延迟。",
    "安全性：登录 Token、关系指令和消息正文会经过该链路。私网、VPN 或专线可减少悟空 WebSocket 直接暴露公网的攻击面。",
    "可控性：固定路由、固定源 IP 和安全组便于限流、审计、故障定位与快速撤销访问。",
    "不影响叨叨：Bridge 只调用现有接口并建立普通 PC(2) 连接，网络接入失败时可以关闭 Bridge，不需要回滚叨叨代码。",
]:
    add_bullet(doc, text)

add_heading(doc, "6. 网络质量与代理配置建议", 1)
add_table(
    doc,
    ["项目", "首期建议值", "说明"],
    [
        ("连接数", "至少允许 1,000 条来自 Bridge 的并发 TCP/WebSocket", "覆盖 500 用户、主备切换和测试余量"),
        ("WebSocket 空闲超时", "不少于 120 秒，建议 300 秒以上", "Bridge 每 30 秒发送心跳，不能由代理在 60 秒左右强制断开"),
        ("TCP keepalive", "允许；不丢弃长连接", "负载均衡和防火墙需保持状态"),
        ("往返时延 RTT", "建议小于 100 ms", "用于支撑端到端 P95 指标"),
        ("丢包率", "建议小于 0.1%", "避免 WebSocket 频繁重连"),
        ("文本阶段带宽", "建议至少 10 Mbps", "文本和控制帧占用较低"),
        ("媒体阶段带宽", "按并发和文件大小评估，建议预留 50 Mbps 以上", "阶段 4 前再压测确认"),
        ("DNS / SNI", "提供固定内网 FQDN 或 LB 地址", "不要依赖容器内仅本机可解析的 `master` 名称"),
    ],
    [3.2, 5.1, 8.0],
)

add_heading(doc, "7. 叨叨研发需要提供的信息", 1)
for text in [
    "叨叨业务 HTTP 的真实内网基础地址、Host/SNI、完整路径以及是否要求 `token` 请求头。",
    "悟空 WebSocket 的真实 URL，例如 `ws://<内网域名>:15200/<path>`，以及反向代理路径和 TLS 要求。",
    "PC(2) JSON-RPC 的 `connect`、`recv`、`recvack`、`send`、`ping` 脱敏真实帧和成功 reasonCode。",
    "测试租户、邀请码、用户名注册开关、限频、Token 有效期和错误码说明。",
    "Bridge 可使用的测试账号，以及允许测试的两个叨叨用户/会话。",
    "Bridge-1/2 应使用的内网网段、路由、DNS 和安全组联系人。",
    "HTTP 上传大小、消息同步接口限制、悟空连接数限制和网关超时配置。",
]:
    add_bullet(doc, text)

add_heading(doc, "8. 安全要求", 1)
for text in [
    "叨叨侧仅放行 Bridge 固定源 IP；环境切换时分别维护测试和生产白名单。",
    "跨公网时 HTTP 必须使用 HTTPS，WebSocket 必须使用 WSS；私网明文 WS 只能运行在受控 VPC/VPN 内。",
    "Token、密码、邀请码和腾讯密钥不写入代码、普通日志或对接文档，通过环境变量或密钥文件注入。",
    "Bridge 使用独立 MySQL 与 Redis，不连接叨叨数据库，也不要求叨叨提供数据库账号。",
    "PHP 到 Bridge 使用 HMAC、Timestamp、Nonce、HTTPS 和 IP 白名单；该链路不需要叨叨开放端口。",
    "出现异常可立即撤销 Bridge 源 IP 的路由/白名单，叨叨原有 App 和服务继续运行。",
]:
    add_bullet(doc, text)

add_heading(doc, "9. 联调与验收步骤", 1)
steps = [
    "网络层：从 Bridge-1/2 解析并访问叨叨 HTTP 内网域名，确认 TCP 建连和返回码。",
    "WebSocket 层：建立测试 PC(2) 连接，保持 30 分钟，验证 ping/pong 无周期性断线。",
    "协议层：完成一次登录、connect、recv、recvack、send，并保存脱敏帧到契约测试。",
    "断线测试：主动中断路由 60 秒后恢复，验证指数重连和同步补偿。",
    "主备测试：终止主 Bridge 租约，确认备用实例 20 秒内接管且没有重复投递。",
    "容量测试：逐步提升至 500 条 PC(2) 连接，观察悟空连接数、CPU、内存、延迟和错误码。",
    "稳定性测试：连续运行 72 小时，对账差异为零，死信为零或每条有明确原因。",
]
for step in steps:
    add_number(doc, step)

add_heading(doc, "10. 验收标准", 1)
add_table(
    doc,
    ["验收项", "通过标准"],
    [
        ("可达性", "Bridge-1/2 均可访问确认后的 HTTP 与 WebSocket 入口"),
        ("长连接", "500 条连接压测期间无固定周期批量断线"),
        ("切换", "主实例失去租约后关闭代理，备用实例 20 秒内接管"),
        ("消息", "文本双向无重复、无回环；断线后可补齐"),
        ("延迟", "叨叨 → 华为 P95 < 1 秒；华为 → 叨叨 P95 < 1.5 秒"),
        ("边界", "叨叨服务端与悟空 IM 代码、配置校验和与联调前一致"),
        ("回滚", "撤销 Bridge 网络访问后，叨叨原有业务不受影响"),
    ],
    [4.2, 12.5],
)

add_heading(doc, "11. 对接信息回填表", 1)
add_table(
    doc,
    ["项目", "叨叨侧回填"],
    [
        ("业务 HTTP 内网 URL", ""),
        ("悟空 WebSocket 内网 URL", ""),
        ("公网测试 HTTP URL（如有）", ""),
        ("公网测试 WSS URL（如有）", ""),
        ("Host / SNI / WebSocket path", ""),
        ("Bridge 允许的源网段", ""),
        ("VPN / 对等连接联系人", ""),
        ("测试租户与邀请码", ""),
        ("测试账号", ""),
        ("连接数与限频", ""),
        ("维护窗口", ""),
    ],
    [6.0, 10.7],
)

add_callout(
    doc,
    "最终确认",
    "本对接仅增加网络访问关系。叨叨服务端、悟空 IM、叨叨数据库和叨叨客户端均不需要修改。若叨叨能够稳定提供公网 HTTPS + WSS，并满足白名单、长连接和容量要求，也可以不用私网专线；正式上线前仍需完成同样的协议、断线、容量和 72 小时稳定性验收。",
    LIGHT_GREEN,
    GREEN,
)

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
